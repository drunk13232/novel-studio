"""Read-only manuscript assembly; python-docx and external Pandoc own formats."""
import html, json, os, re, shutil, subprocess, sys, tempfile, uuid
from pathlib import Path
from fastapi import APIRouter
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict, Field
from app.errors import ConsoleError
from app.library.storage import safe, read, write, digest
from app.studio import body_counts
from app.projects.writing_config import writing,story_row

def heading(c):
    return c['title'] if c.get('kind') in {'story','anthology','paratext'} else f'第 {c["number"]} 章　{c["title"]}'

PRESETS={'web':{'font':'Microsoft YaHei','size':12,'spacing':6,'indent':2,'leading':1.6,'margin':2,'page_break':False},'literary':{'font':'SimSun','size':12,'spacing':10,'indent':2,'leading':1.8,'margin':2.5,'page_break':True},'submission':{'font':'SimSun','size':12,'spacing':0,'indent':2,'leading':1.5,'margin':2.5,'page_break':True}}
def prose(text):
    return re.sub(r'^\ufeff?(?:#{1,6}\s+[^\n]+|第[一二三四五六七八九十百千零〇\d]+章[^\n]*)(?:\n+|$)','',text.replace('\r\n','\n'),count=1)
class ExportRequest(BaseModel):
    model_config=ConfigDict(extra='forbid')
    format:str=Field(default='txt',pattern='^(txt|docx|epub)$')
    author:str=Field(default='',max_length=100)
    description:str=Field(default='',max_length=4000)
    volume:str=''
    chapters:list[str]=Field(default_factory=list,max_length=2000)
    first:int|None=Field(default=None,ge=1)
    last:int|None=Field(default=None,ge=1)
    drafts:bool=False
    bom:bool=False
    cover:str='official'
    preset:str=Field(default='literary',pattern='^(web|literary|submission)$')
    font:str=Field(default='SimSun',max_length=80)
    size:float=Field(default=12,ge=8,le=28)
    spacing:float=Field(default=10,ge=0,le=30)
    indent:float=Field(default=2,ge=0,le=4)
    leading:float=Field(default=1.8,ge=1,le=3)
    margin:float=Field(default=2.5,ge=1,le=5)
    page_break:bool=True
    toc:bool=True
    title_style:str=Field(default='center',pattern='^(center|left)$')
    directory:str=''
    confirm_warnings:bool=False

class PrintRequest(ExportRequest):
    # Reuse selection, author, fonts and paragraph controls. PDF is one artifact,
    # not a second assembler or a browser-only approximation.
    format:str=Field(default='pdf',pattern='^pdf$')
    paper_size:str=Field(default='A5',pattern='^(A5|POCKET32)$')
    gutter_mm:float=Field(default=3,ge=0,le=12)
    running_header:bool=True
    page_numbers:bool=True
    request_id:str|None=Field(default=None,pattern=r'^[0-9a-f]{32}$')

class BookCenter:
    def __init__(self,registry,planning,assets,config):
        self.registry=registry;self.planning=planning;self.assets=assets;self.config=config
        from app.print_preview import PrintPreview
        self.prints=PrintPreview(self,heading)
    def pandoc(self):
        cfg=read(self.config/'book-service.json',{});configured=cfg.get('pandoc','')
        if configured:return Path(configured) if Path(configured).is_file() else None
        found=shutil.which('pandoc')
        if found:return Path(found)
        local=self.registry.engine/'tools/pandoc/pandoc.exe'
        return local if local.is_file() else None
    def capabilities(self):
        path=self.pandoc();version=None;error=None
        if path:
            try:version=subprocess.run([str(path),'--version'],capture_output=True,text=True,timeout=8,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0)).stdout.splitlines()[0]
            except Exception as e:error=str(e)
        from app.print_preview import capabilities as print_capabilities
        return {**print_capabilities(),'pandoc_path':str(path) if path else '', 'pandoc_version':version,'epub_available':bool(version),'error':error,'presets':PRESETS,'default_directory':str(self.registry.engine.parent/'NovelStudio Exports')}
    def catalog(self,p,include_auxiliary=True,include_text_counts=True):
        ws=self.registry.workspace(p);manifest=self.registry.get(p);cfg=writing(manifest);volumes=self.planning.rows(p,'volume');cards=self.planning.rows(p,'chapter');numbers=[c['number'] for c in cards]
        if len(numbers)!=len(set(numbers)):raise ConsoleError('章号重复，无法对应正式正文；请先通过规划界面处理',409)
        ranks={v['id']:i for i,v in enumerate(sorted(volumes,key=lambda v:v['number']))};pref=read(safe(ws.root,'metadata/studio/board.json'),{}).get('chapters',[])
        cards.sort(key=lambda c:(ranks.get(c['volume'],999),pref.index(c['id']) if c['id'] in pref else len(pref)+c['order']))
        manual=read(safe(ws.root,'metadata/manual-chapters.json'),{})
        for folder in safe(ws.root,'chapters').glob('chapter-*'):
            if re.fullmatch(r'chapter-\d+',folder.name) and int(folder.name.split('-')[1]) not in numbers:
                n=int(folder.name.split('-')[1]);m=manual.get(str(n),{});cards.append({'id':'manual_'+digest(('chapter-'+str(n)).encode())[:24],'number':n,'title':m.get('title') or f'第 {n} 章','volume':m.get('volume_id',''),'status':'MANUAL','order':m.get('order',n)})
        cards.sort(key=lambda c:(ranks.get(c['volume'],999),pref.index(c['id']) if c['id'] in pref else len(pref)+c['order']))
        rows=[]
        for c in cards:
            folder=safe(ws.root,f'chapters/chapter-{c["number"]}');final=next((f for f in ['final.txt','final.md'] if (folder/f).is_file()),None);draft=next((f for f in ['draft.txt','draft.md'] if (folder/f).is_file()),None)
            path=f'chapters/chapter-{c["number"]}/{final or draft}' if final or draft else None
            text=prose(ws.content(path)['content']) if path and include_text_counts else ''
            rows.append({'id':c['id'],'number':c['number'],'missing_title':not c['title'],'title':c['title'] or '未命名章节','volume':c['volume'],'status':'FINAL' if final else 'DRAFT' if draft else 'UNWRITTEN','final':f'chapters/chapter-{c["number"]}/{final}' if final else None,'draft':f'chapters/chapter-{c["number"]}/{draft}' if draft else None,**body_counts(text)})
        story=story_row(ws,manifest,include_text_counts=include_text_counts)
        if story:rows.insert(0,{**story,'volume':'','missing_title':False,'draft':story['path'] if story['status']!='FINAL' else story['draft'],'final':story['final_path']})
        if cfg['form']=='anthology':
            rows=[];volumes=[]
            for i,id in enumerate(read(safe(ws.root,'metadata/anthology.json'),{}).get('entries',[]),1):
                try:
                    source=self.registry.get(id)
                    if id==p or writing(source)['form']!='short':raise ConsoleError('篇目不再属于短篇',409)
                    source_catalog=self.catalog(id,include_auxiliary=include_auxiliary,include_text_counts=include_text_counts);items=source_catalog['chapters']
                    rows.append({'id':id,'source_project':id,'kind':'anthology','number':i,'title':read(safe(ws.root,'metadata/anthology.json'),{}).get('display_titles',{}).get(id) or source['title'],'volume':'','missing_title':False,'status':'FINAL' if items and all(c['final'] for c in items) else 'DRAFT','final':None,'draft':None,'han':sum(c['han'] for c in items),'characters':sum(c['characters'] for c in items),'available':True})
                except ConsoleError:rows.append({'id':id,'source_project':id,'kind':'anthology','number':i,'title':'不可用篇目','volume':'','missing_title':False,'status':'UNAVAILABLE','final':None,'draft':None,'han':0,'characters':0,'available':False})
        return {'title':manifest['title'],'description':read(safe(ws.root,'metadata/anthology.json'),{}).get('description',''),'writing':cfg,'volumes':sorted(volumes,key=lambda v:v['number']),'chapters':rows,'preferences':read(safe(ws.root,'metadata/studio/export.json'),{}),'capabilities':self.capabilities() if include_auxiliary else {},'assets':self.assets.list(p) if include_auxiliary else []}
    def assemble(self,p,b):
        catalog=self.catalog(p);ws=self.registry.workspace(p);selected=[];warnings=[]
        if b.first and b.last and b.first>b.last:raise ConsoleError('起始章节不能大于结束章节',422)
        if b.chapters and not set(b.chapters)<={c['id'] for c in catalog['chapters']}:raise ConsoleError('章节稳定 ID 已变化',409)
        for c in catalog['chapters']:
            if b.volume and c['volume']!=b.volume or b.chapters and c['id'] not in b.chapters or b.first and c['number']<b.first or b.last and c['number']>b.last:continue
            if c.get('kind')=='anthology':
                if not c['available']:raise ConsoleError('合集存在不可用篇目，请先修复引用',409)
                sub=b.model_copy(update={'volume':'','chapters':[],'first':None,'last':None,'cover':'none'})
                _,parts,notes,_=self.assemble(c['source_project'],sub)
                warnings.extend(c['title']+'：'+note for note in notes)
                if not parts or any('未完成，未导出' in note or '正文为空' in note for note in notes):raise ConsoleError('合集篇目正文缺失或为空，请修复后导出：'+c['title'],409)
                text='\n\n'.join((heading(item)+'\n\n' if item.get('kind')!='story' else '')+item['text'] for item in parts)
                selected.append({**c,'text':text,'volume_title':''});continue
            path=c['final'] or (c['draft'] if b.drafts else None)
            label=heading(c)
            if not path:warnings.append(label+'未完成，未导出');continue
            document=ws.content(path);raw=document['content'];text=raw if c.get('kind')=='story' else prose(raw)
            selected.append({**c,'text':text,'document_revision':document['etag'],'volume_title':next((v['title'] for v in catalog['volumes'] if v['id']==c['volume']), '')})
            if not text.strip():warnings.append(label+'正文为空')
            if '\ufffd' in text or '\x00' in text:warnings.append(label+'可能含乱码或控制字符')
            if c['missing_title']:warnings.append(label+'缺少标题')
            if not c['final']:warnings.append(label+'使用草稿')
        if not selected:raise ConsoleError('范围内没有可导出的正文；草稿需要主动开启',422)
        if [c['number'] for c in selected]!=sorted(c['number'] for c in selected):warnings.append('显示顺序与章号顺序不同，请核对目录')
        if catalog['writing']['form']=='anthology':
            extra=read(safe(ws.root,'metadata/anthology.json'),{})
            for key,title,at in [('introduction','前言',0),('afterword','后记',len(selected)+1)]:
                if extra.get(key):selected.insert(at,{'id':key,'kind':'paratext','number':0,'title':title,'volume':'','volume_title':'','text':extra[key],**body_counts(extra[key])})
        cover=None
        if b.cover!='none' and b.format!='txt':
            asset=read(safe(ws.root,'covers/official.json'),{}).get('asset_id') if b.cover=='official' else b.cover
            if asset:
                try:cover=self.assets.path(p,asset)
                except ConsoleError:warnings.append('选定封面无法读取')
            else:warnings.append('没有正式封面，将不使用封面')
        return catalog,selected,warnings,cover
    def preview(self,p,b):
        catalog,rows,warnings,cover=self.assemble(p,b)
        return {'title':catalog['title'],'chapters':rows,'warnings':warnings,'han':sum(c['han'] for c in rows),'cover_url':f'/api/projects/{p}/covers/assets/{read(safe(self.registry.workspace(p).root,"covers/official.json"),{}).get("asset_id")}' if b.cover=='official' and cover else f'/api/projects/{p}/assets/{b.cover}/file' if cover else None}
    def export(self,p,b):
        catalog,rows,warnings,cover=self.assemble(p,b)
        if warnings and not b.confirm_warnings:raise ConsoleError('请先预览并确认导出警告',409)
        directory=Path(b.directory).expanduser() if b.directory else self.registry.engine.parent/'NovelStudio Exports'
        directory=directory.resolve()
        if directory.is_relative_to(self.registry.root.resolve()) or directory.is_relative_to(self.registry.engine.resolve()):raise ConsoleError('输出必须在小说工程目录之外',403)
        directory.mkdir(parents=True,exist_ok=True);name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',catalog['title']).strip('. ')[:70] or '小说';id=uuid.uuid4().hex;target=directory/f'{name}-{id[:8]}.{b.format}'
        try:
            if b.format=='txt':
                parts=[catalog['title'],b.author,b.description]
                if b.toc and catalog['writing']['form']!='short':parts+=['目录',*[heading(c) for c in rows]]
                previous=None
                for c in rows:
                    if c['volume'] and c['volume']!=previous:parts.append(c['volume_title']);previous=c['volume']
                    parts.extend([heading(c),c['text']])
                target.write_text('\n\n'.join(parts),encoding='utf-8-sig' if b.bom else 'utf8')
            elif b.format=='docx':self.docx(target,catalog['title'],rows,b,cover)
            else:self.epub(target,catalog['title'],rows,b,cover)
        except Exception as e:
            target.unlink(missing_ok=True)
            if isinstance(e,ConsoleError):raise
            raise ConsoleError('导出失败：'+str(e),422)
        record={'id':id,'path':str(target),'filename':target.name,'format':b.format,'sha256':digest(target.read_bytes()),'warnings':warnings}
        write(self.config/'exports'/p/(id+'.json'),record)
        return {**record,'download_url':f'/api/projects/{p}/book/exports/{id}'}
    def docx(self,path,title,rows,b,cover):
        from docx import Document
        from docx.shared import Pt,Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document();section=doc.sections[0]
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(b.margin)
        style=doc.styles['Normal'];style.font.name=b.font;style.font.size=Pt(b.size);style.element.rPr.rFonts.set(qn('w:eastAsia'),b.font)
        style.paragraph_format.line_spacing=b.leading;style.paragraph_format.space_after=Pt(b.spacing);style.paragraph_format.first_line_indent=Pt(b.size*b.indent)
        if cover:
            from PIL import Image
            with Image.open(cover) as im:w,h=im.size
            maxw=section.page_width-section.left_margin-section.right_margin;maxh=Cm(18)
            doc.add_picture(str(cover),width=min(maxw,int(maxh*w/h)))
        doc.add_heading(title,0);doc.add_paragraph(b.author);doc.add_paragraph(b.description)
        if b.toc:
            doc.add_page_break();doc.add_heading('目录',1)
            p=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(field)
            for c in rows:doc.add_paragraph((c['volume_title']+' · ' if c['volume_title'] else '')+heading(c),style='TOC Heading')
            update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
        previous=None
        for c in rows:
            if c['volume'] and c['volume']!=previous:doc.add_page_break();doc.add_heading(c['volume_title'],1);previous=c['volume']
            elif b.page_break:doc.add_page_break()
            h=doc.add_heading(heading(c),2);h.alignment=WD_ALIGN_PARAGRAPH.CENTER if b.title_style=='center' else WD_ALIGN_PARAGRAPH.LEFT
            for line in c['text'].splitlines():
                if line.strip():doc.add_paragraph(line)
        section.header.paragraphs[0].text=title
        page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');section.footer.paragraphs[0]._p.append(page)
        doc.save(path)
    def epub(self,path,title,rows,b,cover):
        pandoc=self.pandoc()
        if not pandoc:raise ConsoleError('Pandoc 不可用，EPUB 导出已禁用',503)
        with tempfile.TemporaryDirectory(prefix='novel-book-',dir=self.registry.engine/'web/runtime') as tmp:
            folder=Path(tmp);parts=[];previous=None
            for c in rows:
                if c['volume'] and c['volume']!=previous:parts.append('<h1>'+html.escape(c['volume_title'])+'</h1>');previous=c['volume']
                parts.append('<h2>'+html.escape(heading(c))+'</h2>');parts.extend('<p>'+html.escape(line)+'</p>' for line in c['text'].splitlines() if line.strip())
            source=folder/'book.html';source.write_text('<html lang="zh-CN"><head><meta charset="utf-8"></head><body>'+''.join(parts)+'</body></html>','utf-8')
            import yaml
            meta=folder/'metadata.yaml';meta.write_text(yaml.safe_dump({'title':title,'author':[b.author] if b.author else [],'description':b.description,'lang':'zh-CN'},allow_unicode=True),'utf-8')
            css=folder/'book.css';font=b.font.replace('"','').replace("'",'').replace(';','');css.write_text(f'body{{font-family:"{font}",serif;font-size:{b.size}pt;line-height:{b.leading};margin:5%;}}p{{text-indent:{b.indent}em;margin-bottom:{b.spacing}pt}}h1,h2{{text-align:{b.title_style};}}img{{max-width:100%;height:auto}}','utf-8')
            cmd=[str(pandoc),str(source),'--from=html','--to=epub3','--output='+str(path),'--metadata-file='+str(meta),'--css='+str(css),'--split-level=2','--standalone']
            if b.toc:cmd+=['--toc','--toc-depth=2']
            if cover:cmd+=['--epub-cover-image='+str(cover)]
            result=subprocess.run(cmd,capture_output=True,text=True,encoding='utf8',timeout=120,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            if result.returncode:raise ConsoleError('Pandoc EPUB 转换失败：'+result.stderr[:1200],422)

def book_routes(service):
    router=APIRouter(prefix='/api/projects/{p}/book')
    @router.get('')
    def catalog(p:str):return service.catalog(p)
    @router.put('/service')
    def configure(p:str,b:dict):
        service.registry.get(p);path=str(b.get('pandoc',''))
        if path and (not Path(path).is_file() or Path(path).name.lower()!='pandoc.exe'):raise ConsoleError('请选择已有 pandoc.exe',422)
        write(service.config/'book-service.json',{'pandoc':path});return service.capabilities()
    @router.put('/preferences')
    def preferences(p:str,b:ExportRequest):
        service.registry.get(p);write(safe(service.registry.workspace(p).root,'metadata/studio/export.json'),b.model_dump());return {'saved':True}
    @router.post('/preview')
    def preview(p:str,b:ExportRequest):return service.preview(p,b)
    @router.post('/print-preview')
    def print_preview(p:str,b:PrintRequest):return service.prints.preview(p,b)
    @router.post('/print-preview/{request_id}/cancel')
    def cancel_print(p:str,request_id:str):return service.prints.cancel(p,request_id)
    @router.get('/print/{id}/source-status')
    def print_source_status(p:str,id:str):return service.prints.source_status(p,id)
    @router.get('/print/{id}/file')
    def print_file(p:str,id:str):
        record,path=service.prints.get(p,id)
        return FileResponse(path,media_type='application/pdf',filename=record['filename'],content_disposition_type='inline')
    @router.post('/print/{id}/export')
    def print_export(p:str,id:str,b:dict):
        directory=b.get('directory','')
        if not isinstance(directory,str) or len(directory)>2000:raise ConsoleError('请选择有效的 PDF 输出目录',422)
        allow_snapshot=b.get('allow_snapshot',False)
        if not isinstance(allow_snapshot,bool):raise ConsoleError('请选择是否导出旧快照',422)
        return service.prints.export(p,id,directory,allow_snapshot)
    @router.post('/export')
    def export(p:str,b:ExportRequest):return service.export(p,b)
    @router.get('/exports/{id}')
    def file(p:str,id:str):
        service.registry.get(p)
        if not re.fullmatch('[0-9a-f]{32}',id):raise ConsoleError('导出记录不存在',404)
        row=read(service.config/'exports'/p/(id+'.json'))
        if not row or not Path(row['path']).is_file() or digest(Path(row['path']).read_bytes())!=row['sha256']:raise ConsoleError('导出文件缺失或被外部修改',409)
        return FileResponse(row['path'],filename=row['filename'])
    return router
